"""GET /api/v1/reports/risks: report table and the xlsx export."""
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from src.db import async_session_factory
from src.main import app
from src.models.audit import AuditLogEntry
from src.models.auth import User, UserPermission, UserScope, UserScopeFacility
from src.models.hierarchy import Facility
from src.models.risk import Risk, RiskDecision
from src.services.auth_service import hash_password
from src.services.demo_seed import seed_demo_users
from src.services.risk_report import COLUMNS

PERIOD = {"from": "2045-03-01T00:00:00Z", "to": "2045-03-31T23:59:59Z"}
SCOPED_USER = "usr_report_scoped"


def _risk(risk_id: str, facility_id: str, risk_type: str, probability: float, level: str, as_of: datetime, status: str) -> Risk:
    return Risk(
        id=risk_id, forecast_id=risk_id, risk_type=risk_type, target_type="facility", target_id=facility_id,
        facility_id=facility_id, as_of=as_of, lead_min_hours=0.0, horizon_hours=72.0,
        prediction_window_start=as_of, prediction_window_end=as_of + timedelta(hours=72), probability=probability,
        threshold=0.6, risk_level=level, priority_score=50.0, decision_status=status, sla_due_at=as_of + timedelta(hours=24),
        data_health="fresh", model="report-test", top_factors=[], recommendation="r", version=2,
        created_at=as_of, updated_at=as_of,
    )


async def _client_token(client: AsyncClient, username: str, password: str) -> str:
    return (await client.post("/api/v1/auth/login", json={"username": username, "password": password})).json()["token"]


@pytest.fixture
async def seeded() -> dict:
    tag = uuid.uuid4().hex[:6]
    ids = {"rejected": f"risk_rpt_{tag}_1", "acknowledged": f"risk_rpt_{tag}_2", "outside": f"risk_rpt_{tag}_3"}
    fac_a, fac_b = f"fac_rpt_{tag}_a", f"fac_rpt_{tag}_b"
    async with async_session_factory() as session:
        await seed_demo_users(session)
        session.add(Facility(id=fac_a, display_name="Объект А; корпус \"1\"", facility_type="test", district_id=None))
        session.add(Facility(id=fac_b, display_name="Объект Б", facility_type="test", district_id=None))
        await session.flush()
        session.add(_risk(ids["rejected"], fac_a, "fire", 0.72, "medium", datetime(2045, 3, 1, 10, 0, tzinfo=timezone.utc), "rejected"))
        session.add(_risk(ids["acknowledged"], fac_b, "flooding", 0.9, "high", datetime(2045, 3, 2, 12, 0, tzinfo=timezone.utc), "acknowledged"))
        session.add(_risk(ids["outside"], fac_a, "fire", 0.5, "medium", datetime(2045, 4, 15, 9, 0, tzinfo=timezone.utc), "open"))
        await session.flush()
        session.add(RiskDecision(id=f"dec_{tag}_1", risk_id=ids["rejected"], user_id="usr_manager", decision="acknowledged",
                                 reason_code=None, comment="сначала подтвердили", decided_at=datetime(2045, 3, 1, 11, tzinfo=timezone.utc)))
        session.add(RiskDecision(id=f"dec_{tag}_2", risk_id=ids["rejected"], user_id="usr_manager", decision="rejected",
                                 reason_code="false_alarm", comment="пыль, строка\nвторая", decided_at=datetime(2045, 3, 1, 12, tzinfo=timezone.utc)))
        if await session.get(User, SCOPED_USER) is None:
            session.add(User(id=SCOPED_USER, username=SCOPED_USER, password_hash=hash_password("password123"), display_name="r", role="test"))
            session.add(UserPermission(user_id=SCOPED_USER, permission="report.export"))
            session.add(UserScope(user_id=SCOPED_USER, scope_type="assigned_facilities"))
        await session.flush()
        session.add(UserScopeFacility(user_id=SCOPED_USER, facility_id=fac_a))
        await session.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        token = await _client_token(client, "manager", "manager123")
        created = await client.post(
            "/api/v1/work-orders", headers={"Authorization": f"Bearer {token}"},
            json={
                "mode": "draft", "source_risk_id": ids["acknowledged"], "facility_id": fac_b, "target_entity_type": "facility",
                "target_entity_id": fac_b, "work_type": "inspection", "priority": "high",
                "due_at": datetime(2052, 1, 1, tzinfo=timezone.utc).isoformat(), "description": "report", "comment": None,
            },
        )
    assert created.status_code == 201, created.text
    return {"ids": ids, "fac_a": fac_a, "fac_b": fac_b, "work_order": created.json()["display_number"]}


async def _export(username: str, password: str, params: dict):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        token = await _client_token(client, username, password)
        return await client.get("/api/v1/reports/risks", headers={"Authorization": f"Bearer {token}"}, params=params)


def _xlsx_rows(content: bytes) -> tuple[list, dict[str, list]]:
    sheet = load_workbook(io.BytesIO(content)).active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    return rows[0], {row[0]: row for row in rows[1:]}


@pytest.mark.asyncio
async def test_xlsx_content_is_typed_and_complete(seeded: dict) -> None:
    response = await _export("manager", "manager123", PERIOD)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="risks_2045-03-01_2045-03-31.xlsx"' in response.headers["content-disposition"]

    header, rows = _xlsx_rows(response.content)
    assert header == COLUMNS
    ids = seeded["ids"]
    assert ids["outside"] not in rows, "risks outside the period must not be exported"

    rejected = dict(zip(COLUMNS, rows[ids["rejected"]]))
    assert rejected["Объект"] == seeded["fac_a"]
    assert rejected["Название объекта"] == "Объект А; корпус \"1\""
    assert rejected["Тип риска"] == "Пожар"
    assert rejected["Вероятность"] == pytest.approx(0.72) and isinstance(rejected["Вероятность"], float)
    assert rejected["Уровень"] == "Средний"
    assert rejected["Время прогноза (МСК)"] == datetime(2045, 3, 1, 13, 0), "10:00 UTC is 13:00 MSK, stored as a real date"
    assert rejected["Решение"] == "Отклонён"
    assert rejected["Причина отклонения"] == "Ложное срабатывание", "the latest decision's reason, by display name"
    assert rejected["Комментарий"] == "пыль, строка\nвторая"
    assert rejected["Заявка"] in (None, "")

    acknowledged = dict(zip(COLUMNS, rows[ids["acknowledged"]]))
    assert (acknowledged["Тип риска"], acknowledged["Уровень"], acknowledged["Решение"]) == ("Подтопление", "Высокий", "Подтверждён")
    assert acknowledged["Заявка"] == seeded["work_order"]


@pytest.mark.asyncio
async def test_default_format_is_xlsx(seeded: dict) -> None:
    response = await _export("manager", "manager123", {"from": PERIOD["from"]})
    assert response.status_code == 200
    assert 'filename="risks_2045-03-01_all.xlsx"' in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_scope_limits_the_export(seeded: dict) -> None:
    response = await _export(SCOPED_USER, "password123", PERIOD)
    assert response.status_code == 200, response.text
    _, rows = _xlsx_rows(response.content)
    assert seeded["ids"]["rejected"] in rows
    assert seeded["ids"]["acknowledged"] not in rows, "facility outside the caller's scope"


@pytest.mark.asyncio
async def test_bad_format_and_bad_date_are_400(seeded: dict) -> None:
    assert (await _export("manager", "manager123", {**PERIOD, "format": "pdf"})).status_code == 400
    assert (await _export("manager", "manager123", {"from": "yesterday"})).status_code == 400


@pytest.mark.asyncio
async def test_role_without_report_export_is_denied(seeded: dict) -> None:
    # The coordinator manages repairs but has no report.export in the demo role model.
    assert (await _export("coordinator", "coordinator123", PERIOD)).status_code == 403


@pytest.mark.asyncio
async def test_facility_dispatcher_exports_only_own_facility(seeded: dict) -> None:
    # The facility dispatcher holds report.export but sees only its assigned facility.
    response = await _export("dispatcher", "dispatcher123", PERIOD)
    assert response.status_code == 200, response.text
    _, rows = _xlsx_rows(response.content)
    assert not set(seeded["ids"].values()) & set(rows), "facilities outside the dispatcher's scope must not be exported"


@pytest.mark.asyncio
async def test_exports_are_journaled_including_denials(seeded: dict) -> None:
    ok = await _export("manager", "manager123", {**PERIOD, "format": "xlsx"})
    denied = await _export("coordinator", "coordinator123", PERIOD)
    async with async_session_factory() as session:
        entries = {
            e.trace_id: e
            for e in (
                await session.execute(
                    select(AuditLogEntry).where(AuditLogEntry.trace_id.in_([ok.headers["X-Trace-Id"], denied.headers["X-Trace-Id"]]))
                )
            ).scalars()
        }
    exported = entries[ok.headers["X-Trace-Id"]]
    assert exported.action == "GET /api/v1/reports/risks"
    assert (exported.user_id, exported.result) == ("usr_manager", "success")
    assert exported.details["format"] == "xlsx" and exported.details["from"] == PERIOD["from"]
    refused = entries[denied.headers["X-Trace-Id"]]
    assert (refused.user_id, refused.result, refused.status_code) == ("usr_coordinator", "denied", 403)
