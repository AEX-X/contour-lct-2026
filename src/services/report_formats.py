"""Render a ReportTable into a downloadable file in the requested format."""
import csv
import io
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from functools import partial
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from src.services.risk_report import PROBABILITY_COLUMN, TIME_COLUMN, ReportTable

DEFAULT_FORMAT = "xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
# Excel detects UTF-8 only by the byte-order mark; without it Cyrillic turns into mojibake.
UTF8_BOM = "﻿"


@dataclass(frozen=True)
class ReportFile:
    """Rendered bytes with their media type and file extension."""

    content: bytes
    media_type: str
    extension: str


class UnknownReportFormat(ValueError):
    """The requested export format is not supported."""


def render_xlsx(table: ReportTable) -> ReportFile:
    """Excel workbook with typed cells: probability as a number, time as a date."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Риски"
    sheet.append(table.headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in table.rows:
        sheet.append(row)
    for row_cells in sheet.iter_rows(min_row=2):
        row_cells[PROBABILITY_COLUMN].number_format = "0.00"
        row_cells[TIME_COLUMN].number_format = "yyyy-mm-dd hh:mm:ss"
    for index, header in enumerate(table.headers, start=1):
        width = max([len(header)] + [len(str(row[index - 1])) for row in table.rows]) + 2
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(width, 60)
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return ReportFile(buffer.getvalue(), XLSX_MEDIA_TYPE, "xlsx")


def csv_cell(value: Any, decimal_separator: str) -> str:
    """Text for one CSV cell: numbers with the given decimal separator, times as ISO date-time."""
    if value is None:
        return ""
    if isinstance(value, float):
        text = f"{value:.4f}".rstrip("0").rstrip(".")
        return text.replace(".", decimal_separator)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def render_csv(table: ReportTable, *, delimiter: str, decimal_separator: str) -> ReportFile:
    """CSV with a UTF-8 BOM; fields with the delimiter, quotes or line breaks are quoted.

    Args:
        table: The report table.
        delimiter: Column separator.
        decimal_separator: Decimal mark for numbers ("," for Russian Excel, "." otherwise).
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(table.headers)
    for row in table.rows:
        writer.writerow([csv_cell(value, decimal_separator) for value in row])
    return ReportFile((UTF8_BOM + buffer.getvalue()).encode("utf-8"), CSV_MEDIA_TYPE, "csv")


RENDERERS: dict[str, Callable[[ReportTable], ReportFile]] = {
    "xlsx": render_xlsx,
    # Russian-locale Excel: ";" between columns and a decimal comma, opens by double click.
    "csv_semicolon": partial(render_csv, delimiter=";", decimal_separator=","),
    # Classic CSV for English Excel, Google Sheets, LibreOffice and scripts.
    "csv_comma": partial(render_csv, delimiter=",", decimal_separator="."),
}


def render(table: ReportTable, report_format: str) -> ReportFile:
    """Render the table in one of RENDERERS' formats.

    Raises:
        UnknownReportFormat: The format is not in RENDERERS.
    """
    renderer = RENDERERS.get(report_format)
    if renderer is None:
        raise UnknownReportFormat(report_format)
    return renderer(table)
